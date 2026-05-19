"""数据导出 API 测试"""
import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi.testclient import TestClient

from src.db.session import get_db
from src.main import create_app

STUDENT_ID = "stu_export_001"
OTHER_ID = "stu_other_999"


@pytest.fixture
def app():
    _app = create_app()

    async def override_get_db():
        yield AsyncMock()

    _app.dependency_overrides[get_db] = override_get_db
    return _app


@pytest.fixture
def client(app):
    return TestClient(app)


def _make_profile(mastery: float = 0.6, appear: int = 5, error: int = 2):
    p = MagicMock()
    p.mastery_score = mastery
    p.appear_count = appear
    p.error_count = error
    p.review_priority = "medium"
    p.last_reviewed_at = None
    return p


def _make_kp(name: str = "三角函数", subject: str = "数学", grade: str = "八年级"):
    import uuid
    kp = MagicMock()
    kp.id = uuid.UUID("00000000-0000-0000-0000-000000000001")
    kp.name = name
    kp.subject = subject
    kp.grade = grade
    return kp


def _make_question(subject: str = "数学", grade: str = "八年级", raw_text: str = "题目文本"):
    import uuid
    q = MagicMock()
    q.id = uuid.UUID("00000000-0000-0000-0000-000000000002")
    q.raw_text = raw_text
    q.subject = subject
    q.grade = grade
    q.question_type = "填空题"
    q.knowledge_points = ["三角函数"]
    q.difficulty = 3
    q.review_priority = "high"
    q.need_review = True
    q.created_at = None
    return q


class TestExportKnowledgePointsJSON:
    """GET /api/export/student/{student_id}/knowledge-points?format=json"""

    def test_returns_200_json_envelope(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [(_make_profile(), _make_kp())]
        mock_db.execute.return_value = result
        app.dependency_overrides[get_db] = lambda: (x for x in [mock_db])

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/knowledge-points?format=json",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["student_id"] == STUDENT_ID
        assert "total" in data
        assert "knowledge_points" in data

    def test_idor_returns_404(self, client):
        resp = client.get(
            f"/api/export/student/{OTHER_ID}/knowledge-points",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 404

    def test_invalid_format_returns_422(self, client):
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/knowledge-points?format=xml",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 422

    def test_csv_content_type(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = []
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/knowledge-points?format=csv",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]


class TestExportQuestionsJSON:
    """GET /api/export/student/{student_id}/questions?format=json|csv"""

    def test_returns_200_with_envelope(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [_make_question()]
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=json",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["student_id"] == STUDENT_ID
        assert "total" in data
        assert "questions" in data

    def test_idor_returns_404(self, client):
        resp = client.get(
            f"/api/export/student/{OTHER_ID}/questions",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 404

    def test_empty_returns_zero_total(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = []
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=json",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        assert resp.json()["total"] == 0
        assert resp.json()["questions"] == []

    def test_question_item_contains_required_fields(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [_make_question(raw_text="解方程 x+1=0")]
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=json",
            headers={"X-Student-Id": STUDENT_ID},
        )
        q = resp.json()["questions"][0]
        for field in ("question_id", "raw_text", "subject", "grade", "review_priority"):
            assert field in q, f"missing field: {field}"

    def test_csv_content_type(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = []
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=csv",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]

    def test_invalid_format_returns_422(self, client):
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=pdf",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 422


class TestExportKnowledgePointsExcel:
    """GET /api/export/student/{student_id}/knowledge-points?format=excel"""

    def test_returns_200_with_xlsx_content_type(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [(_make_profile(), _make_kp())]
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/knowledge-points?format=excel",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

    def test_excel_response_is_valid_workbook(self, app):
        import io
        import openpyxl

        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [(_make_profile(mastery=0.7), _make_kp("二次函数"))]
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/knowledge-points?format=excel",
            headers={"X-Student-Id": STUDENT_ID},
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.max_row >= 2  # 至少有表头+1行数据


class TestExportQuestionsExcel:
    """GET /api/export/student/{student_id}/questions?format=excel"""

    def test_returns_200_with_xlsx_content_type(self, app):
        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = [_make_question()]
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=excel",
            headers={"X-Student-Id": STUDENT_ID},
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

    def test_excel_empty_returns_header_only(self, app):
        import io
        import openpyxl

        mock_db = AsyncMock()
        result = MagicMock()
        result.all.return_value = []
        mock_db.execute.return_value = result

        async def override():
            yield mock_db

        app.dependency_overrides[get_db] = override
        client = TestClient(app)
        resp = client.get(
            f"/api/export/student/{STUDENT_ID}/questions?format=excel",
            headers={"X-Student-Id": STUDENT_ID},
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.max_row == 1  # 只有表头
